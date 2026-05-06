from collections import OrderedDict, defaultdict
from datetime import datetime, date
from decimal import Decimal, InvalidOperation

import openpyxl
import stripe
from django.conf import settings
from django.contrib.auth import logout
from django.contrib.auth.decorators import login_required, user_passes_test
from django.core.mail import send_mail
from django.db.models import Sum
from django.http import JsonResponse, HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.utils import timezone
from django.views.decorators.csrf import csrf_exempt

from .forms import BlogCommentForm, HousingApplicationForm, FinancialUploadForm
from .models import (
    Property,
    BlogPost,
    HousingApplication,
    Payment,
    FinancialUpload,
    FinancialEntry,
)


LATE_FEE_AMOUNT = Decimal("25.00")


def money(value):
    if value is None:
        return Decimal("0.00")

    if isinstance(value, Decimal):
        return value

    try:
        return Decimal(str(value)).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError, TypeError):
        return Decimal("0.00")


def staff_required(user):
    return user.is_authenticated and (user.is_staff or user.is_superuser)


def classify_financial_entry(sheet_name, category, description):
    text = f"{sheet_name} {category} {description}".lower()

    if "debt" in text or "mortgage" in text or "loan" in text or "principal" in text or "interest" in text:
        return "debt_service"

    if "capex" in text or "capital" in text or "improvement" in text or "renovation" in text:
        return "capital_expense"

    if "rent" in text or "income" in text or "revenue" in text or "deposit" in text:
        return "income"

    if "expense" in text or "repair" in text or "maintenance" in text or "utility" in text or "insurance" in text or "tax" in text:
        return "operating_expense"

    return "other"


def extract_month_year(sheet_name, entry_date=None):
    month_names = {
        "jan": 1,
        "january": 1,
        "feb": 2,
        "february": 2,
        "mar": 3,
        "march": 3,
        "apr": 4,
        "april": 4,
        "may": 5,
        "jun": 6,
        "june": 6,
        "jul": 7,
        "july": 7,
        "aug": 8,
        "august": 8,
        "sep": 9,
        "sept": 9,
        "september": 9,
        "oct": 10,
        "october": 10,
        "nov": 11,
        "november": 11,
        "dec": 12,
        "december": 12,
    }

    if entry_date:
        return entry_date.month, entry_date.year

    lower = str(sheet_name).lower()
    found_month = None

    for name, number in month_names.items():
        if name in lower:
            found_month = number
            break

    found_year = timezone.localdate().year

    for token in lower.replace("-", " ").replace("_", " ").split():
        if token.isdigit() and len(token) == 4:
            found_year = int(token)
            break

    return found_month, found_year


def parse_date_cell(value):
    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    if isinstance(value, str):
        cleaned = value.strip()
        for fmt in ("%m/%d/%Y", "%m-%d-%Y", "%Y-%m-%d", "%m/%d/%y"):
            try:
                return datetime.strptime(cleaned, fmt).date()
            except ValueError:
                pass

    return None


def parse_excel_upload(upload):
    FinancialEntry.objects.filter(upload=upload).delete()

    workbook = openpyxl.load_workbook(upload.file.path, data_only=True)
    created = 0

    skip_sheets = {"receipts"}

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]

        if sheet_name.strip().lower() in skip_sheets:
            continue

        header_row = None
        headers = {}

        for row_index, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 15)), start=1):
            values = [cell.value for cell in row]
            lowered = [str(v).strip().lower() if v is not None else "" for v in values]

            if any("amount" in v for v in lowered) or any("paid" in v for v in lowered) or any("expense" in v for v in lowered):
                header_row = row_index
                for idx, value in enumerate(lowered):
                    if value:
                        headers[value] = idx
                break

        if not header_row:
            continue

        for row_number, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            if not row or not any(row):
                continue

            row_values = list(row)

            def get_by_keywords(keywords):
                for header, index in headers.items():
                    if any(keyword in header for keyword in keywords):
                        if index < len(row_values):
                            return row_values[index]
                return None

            raw_date = get_by_keywords(["date"])
            raw_category = get_by_keywords(["category", "type", "account"])
            raw_description = get_by_keywords(["description", "memo", "vendor", "name", "resident"])
            raw_amount = get_by_keywords(["amount", "paid", "expense", "payment", "debit", "credit"])

            if raw_amount is None:
                numeric_cells = [value for value in row_values if isinstance(value, (int, float, Decimal))]
                if numeric_cells:
                    raw_amount = numeric_cells[-1]

            amount = money(raw_amount)

            if amount == 0:
                continue

            entry_date = parse_date_cell(raw_date)
            month, year = extract_month_year(sheet_name, entry_date)

            category = str(raw_category).strip() if raw_category is not None else ""
            description = str(raw_description).strip() if raw_description is not None else ""

            entry_type = classify_financial_entry(sheet_name, category, description)

            FinancialEntry.objects.create(
                upload=upload,
                property_name="Painted Lady",
                sheet_name=sheet_name,
                row_number=row_number,
                entry_date=entry_date,
                month=month,
                year=year,
                entry_type=entry_type,
                category=category,
                description=description,
                amount=abs(amount),
            )

            created += 1

    upload.parsed_at = timezone.now()
    upload.save(update_fields=["parsed_at"])

    return created


def home(request):
    properties = Property.objects.all()
    posts = BlogPost.objects.all().order_by("-created_at")[:5]

    return render(request, "home.html", {
        "properties": properties,
        "posts": posts,
    })


def logout_view(request):
    logout(request)
    request.session.flush()
    return redirect("login")


def creed(request):
    return render(request, "creed.html")


def apply(request):
    if request.method == "POST":
        form = HousingApplicationForm(request.POST, request.FILES)
        if form.is_valid():
            form.save()
            return redirect("apply_success")
    else:
        form = HousingApplicationForm()

    return render(request, "apply.html", {"form": form})


def apply_success(request):
    return render(request, "apply_success.html")


def enter_invite_code(request):
    return render(request, "enter_invite_code.html")


def signup(request):
    return render(request, "signup.html")


@login_required
def tenant_dashboard(request):
    request.session.set_expiry(0)

    application = HousingApplication.objects.filter(
        email=request.user.email
    ).first()

    payments = []
    if application:
        payments = application.payments.all().order_by("-created_at")

    return render(request, "tenant_dashboard.html", {
        "application": application,
        "payments": payments,
        "stripe_public_key": settings.STRIPE_PUBLIC_KEY,
    })


@login_required
def landlord_dashboard(request):
    applications = HousingApplication.objects.all().order_by("-created_at")
    properties = Property.objects.all()
    payments = Payment.objects.all().order_by("-created_at")[:25]

    return render(request, "landlord_dashboard.html", {
        "applications": applications,
        "properties": properties,
        "payments": payments,
    })


@login_required
@user_passes_test(staff_required)
def payment_log(request):
    completed_payments = (
        Payment.objects
        .filter(status="completed")
        .select_related("application", "application__property")
        .order_by(
            "application__property__name",
            "-created_at",
            "application__space_label",
            "application__full_name",
        )
    )

    grouped = OrderedDict()

    for payment in completed_payments:
        application = payment.application
        property_name = application.property.name if application.property else "No Property"

        month_label = timezone.localtime(payment.created_at).strftime("%B %Y")

        grouped.setdefault(property_name, OrderedDict())
        grouped[property_name].setdefault(month_label, [])
        grouped[property_name][month_label].append(payment)

    payment_log_data = []

    for property_name, months in grouped.items():
        month_data = []

        for month_label, payments in months.items():
            payments_sorted = sorted(
                payments,
                key=lambda p: (
                    (p.application.space_label or p.application.space_type or "").lower(),
                    p.application.full_name.lower(),
                )
            )

            month_data.append({
                "month_label": month_label,
                "payments": payments_sorted,
            })

        payment_log_data.append({
            "property_name": property_name,
            "months": month_data,
        })

    return render(request, "payment_log.html", {
        "payment_log": payment_log_data,
    })


@login_required
@user_passes_test(staff_required)
def rent_roll(request):
    residents = (
        HousingApplication.objects
        .select_related("property")
        .order_by("property__name", "space_label", "full_name")
    )

    rows = []

    for resident in residents:
        completed_payments = resident.payments.filter(status="completed")
        rent_paid = completed_payments.filter(payment_type="rent").aggregate(total=Sum("amount"))["total"] or Decimal("0.00")
        utility_paid = completed_payments.filter(payment_type="utility").aggregate(total=Sum("amount"))["total"] or Decimal("0.00")
        deposit_paid = resident.deposit_paid

        rows.append({
            "property": resident.property.name if resident.property else "No Property",
            "room": resident.space_label or resident.space_type or "—",
            "resident": resident.full_name,
            "monthly_rent": resident.monthly_rent,
            "rent_balance": resident.balance,
            "rent_paid": rent_paid,
            "utility_monthly": resident.utility_monthly,
            "utility_balance": resident.utility_balance,
            "utility_paid": utility_paid,
            "deposit_required": resident.deposit_required,
            "deposit_paid": deposit_paid,
        })

    return render(request, "rent_roll.html", {
        "rows": rows,
    })


@login_required
@user_passes_test(staff_required)
def t12_report(request):
    year = timezone.localdate().year

    financial_entries = FinancialEntry.objects.filter(year=year)

    months = []

    for month_number in range(1, 13):
        online_income = (
            Payment.objects
            .filter(status="completed", created_at__year=year, created_at__month=month_number)
            .aggregate(total=Sum("amount"))["total"] or Decimal("0.00")
        )

        operating_expenses = (
            financial_entries
            .filter(month=month_number, entry_type="operating_expense")
            .aggregate(total=Sum("amount"))["total"] or Decimal("0.00")
        )

        debt_service = (
            financial_entries
            .filter(month=month_number, entry_type="debt_service")
            .aggregate(total=Sum("amount"))["total"] or Decimal("0.00")
        )

        capital_expenses = (
            financial_entries
            .filter(month=month_number, entry_type="capital_expense")
            .aggregate(total=Sum("amount"))["total"] or Decimal("0.00")
        )

        net_operating_income = online_income - operating_expenses
        cash_flow_after_debt = net_operating_income - debt_service

        months.append({
            "month_number": month_number,
            "month_name": date(year, month_number, 1).strftime("%B"),
            "online_income": online_income,
            "operating_expenses": operating_expenses,
            "debt_service": debt_service,
            "capital_expenses": capital_expenses,
            "net_operating_income": net_operating_income,
            "cash_flow_after_debt": cash_flow_after_debt,
        })

    totals = {
        "online_income": sum((m["online_income"] for m in months), Decimal("0.00")),
        "operating_expenses": sum((m["operating_expenses"] for m in months), Decimal("0.00")),
        "debt_service": sum((m["debt_service"] for m in months), Decimal("0.00")),
        "capital_expenses": sum((m["capital_expenses"] for m in months), Decimal("0.00")),
        "net_operating_income": sum((m["net_operating_income"] for m in months), Decimal("0.00")),
        "cash_flow_after_debt": sum((m["cash_flow_after_debt"] for m in months), Decimal("0.00")),
    }

    return render(request, "t12_report.html", {
        "year": year,
        "months": months,
        "totals": totals,
    })


@login_required
@user_passes_test(staff_required)
def financial_upload(request):
    if request.method == "POST":
        form = FinancialUploadForm(request.POST, request.FILES)
        if form.is_valid():
            upload = form.save()
            return redirect("parse_financial_upload", upload_id=upload.id)
    else:
        form = FinancialUploadForm()

    uploads = FinancialUpload.objects.all().order_by("-uploaded_at")

    return render(request, "financial_upload.html", {
        "form": form,
        "uploads": uploads,
    })


@login_required
@user_passes_test(staff_required)
def parse_financial_upload(request, upload_id):
    upload = get_object_or_404(FinancialUpload, id=upload_id)
    created = parse_excel_upload(upload)

    return render(request, "financial_upload_parsed.html", {
        "upload": upload,
        "created": created,
    })


def property_detail(request, pk):
    property_obj = get_object_or_404(Property, pk=pk)
    gallery_images = property_obj.images.all()

    return render(request, "property_detail.html", {
        "property": property_obj,
        "gallery_images": gallery_images,
    })


def blog_detail(request, pk):
    post = get_object_or_404(BlogPost, pk=pk)

    return render(request, "blog_detail.html", {
        "post": post,
    })


def add_blog_comment(request, post_id):
    post = get_object_or_404(BlogPost, id=post_id)

    if request.method == "POST":
        form = BlogCommentForm(request.POST)
        if form.is_valid():
            comment = form.save(commit=False)
            comment.post = post
            comment.approved = False
            comment.save()

    return redirect("home")


def printable_application(request, pk):
    application = get_object_or_404(HousingApplication, pk=pk)

    return render(request, "printable_application.html", {
        "application": application,
    })


@login_required
def create_checkout_session(request, application_id, payment_type="rent"):
    application = get_object_or_404(HousingApplication, id=application_id)

    today = timezone.localdate()
    amount = Decimal("0.00")
    description = ""

    if payment_type == "rent":
        base = application.balance if application.balance > 0 else application.monthly_rent
        amount = base
        description = "Rent Payment"

        if today.day >= 5 and base > 0:
            late_fee_paid = Payment.objects.filter(
                application=application,
                description__icontains="late fee",
                status="completed",
                created_at__month=today.month,
                created_at__year=today.year,
            ).exists()

            if not late_fee_paid:
                amount += LATE_FEE_AMOUNT
                description = "Rent Payment including $25 late fee"

    elif payment_type == "deposit":
        amount = max(application.deposit_required - application.deposit_paid, Decimal("0.00"))
        description = "Deposit Payment"

    elif payment_type == "utility":
        amount = application.utility_balance or application.utility_monthly
        description = "Utilities Payment"

    else:
        return JsonResponse({"error": "Invalid type"}, status=400)

    if amount <= 0:
        return JsonResponse({"error": "No balance due"}, status=400)

    stripe.api_key = settings.STRIPE_SECRET_KEY

    payment = Payment.objects.create(
        application=application,
        payment_type=payment_type,
        description=description,
        amount=amount,
        status="pending",
    )

    session = stripe.checkout.Session.create(
        payment_method_types=["card"],
        mode="payment",
        line_items=[{
            "price_data": {
                "currency": "usd",
                "product_data": {"name": description},
                "unit_amount": int(amount * 100),
            },
            "quantity": 1,
        }],
        success_url=request.build_absolute_uri("/payment-success/"),
        cancel_url=request.build_absolute_uri("/tenant-dashboard/"),
        metadata={"payment_id": str(payment.id)},
    )

    payment.stripe_session_id = session.id
    payment.save()

    return JsonResponse({"id": session.id})


def payment_success(request):
    return render(request, "payment_success.html")


@csrf_exempt
def stripe_webhook(request):
    payload = request.body
    sig = request.META.get("HTTP_STRIPE_SIGNATURE")

    try:
        event = stripe.Webhook.construct_event(payload, sig, settings.STRIPE_WEBHOOK_SECRET)
    except Exception:
        return HttpResponse(status=400)

    if event["type"] != "checkout.session.completed":
        return HttpResponse(status=200)

    session = event["data"]["object"]
    metadata = session["metadata"] if "metadata" in session else {}
    payment_id = metadata["payment_id"] if "payment_id" in metadata else None

    if not payment_id:
        return HttpResponse(status=200)

    payment = Payment.objects.filter(id=payment_id).first()
    if not payment or payment.status == "completed":
        return HttpResponse(status=200)

    application = payment.application
    payment.status = "completed"
    payment.stripe_payment_intent = session["payment_intent"] if "payment_intent" in session else ""
    payment.save()

    if payment.payment_type == "rent":
        reduction = payment.amount
        if "late fee" in payment.description.lower():
            reduction -= LATE_FEE_AMOUNT
        application.balance = max(Decimal("0.00"), application.balance - reduction)

    elif payment.payment_type == "deposit":
        application.deposit_paid = min(
            application.deposit_required,
            application.deposit_paid + payment.amount
        )

    elif payment.payment_type == "utility":
        application.utility_balance = max(
            Decimal("0.00"),
            application.utility_balance - payment.amount
        )

    application.save()

    owner_email = "BowlingLegacyLLC@outlook.com"
    if application.property and application.property.owner_email:
        owner_email = application.property.owner_email

    send_mail(
        subject="Resident Payment Received",
        message=f"""
Payment received from {application.full_name}

Type: {payment.get_payment_type_display()}
Amount: ${payment.amount}

Rent Balance: ${application.balance}
Deposit Paid: ${application.deposit_paid}
Utilities Balance: ${application.utility_balance}
""",
        from_email=getattr(settings, "DEFAULT_FROM_EMAIL", None),
        recipient_list=[owner_email],
        fail_silently=True,
    )

    return HttpResponse(status=200)
